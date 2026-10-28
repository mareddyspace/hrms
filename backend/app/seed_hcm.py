"""Read and validate the official synthetic HCM workbook (backend/seed/TimeZ_500_Employee_Sample_HCM_Data.xlsx).

Check it on its own:  python -m app.seed_hcm
"""

from __future__ import annotations

from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.config import BACKEND_DIR

WORKBOOK = BACKEND_DIR / "seed" / "TimeZ_500_Employee_Sample_HCM_Data.xlsx"

COLUMNS = {
    "Company": "company",
    "Employee ID": "id",
    "First Name": "first_name",
    "Last Name": "last_name",
    "Gender": "gender",
    "Date of Birth": "date_of_birth",
    "Work Email": "email",
    "Country Code": "country_code",
    "Legal Entity": "legal_entity",
    "Location": "location",
    "Business Unit": "business_unit",
    "Department": "department",
    "Job Title": "job_title",
    "Grade": "grade",
    "Employment Type": "employment_type",
    "Employment Status": "employment_status",
    "Joining Date": "joining_date",
    "Currency": "currency",
    "Time Zone": "time_zone",
    "Manager Employee ID": "manager_id",
    "Holiday Calendar": "holiday_calendar",
    "Leave Policy": "leave_policy_name",
    "Payroll Status": "payroll_status",
    "AI Profile Status": "ai_profile_status",
}


class SeedDataError(ValueError):
    pass


def _as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip()[:10])


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _rows(ws) -> list[dict[str, Any]]:  # type: ignore[no-untyped-def]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]


def cost_centre(country_code: str, department: str) -> str:
    code = "".join(w[0] for w in department.replace("&", " ").split()).upper() if " " in department else department[:3].upper()
    return f"CC-{country_code}-{code}"


def load_employees(path: Path = WORKBOOK) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Return (employee rows ready for the employees table, counts). Raises SeedDataError when the data is inconsistent."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        raw = _rows(wb["Employees"])
        country_summary = {r["Country Code"]: int(r["Employees"]) for r in _rows(wb["Country Summary"])}
        dept_summary = {r["Department"]: int(r["Employees"]) for r in _rows(wb["Department Summary"])}
    finally:
        wb.close()

    missing = [c for c in COLUMNS if raw and c not in raw[0]]
    if missing:
        raise SeedDataError(f"Employees sheet is missing columns: {missing}")

    people: list[dict[str, Any]] = []
    for r in raw:
        p: dict[str, Any] = {field: _text(r.get(col)) for col, field in COLUMNS.items()}
        p["date_of_birth"] = _as_date(r.get("Date of Birth"))
        p["joining_date"] = _as_date(r.get("Joining Date"))
        p["country_code"] = (p["country_code"] or "").upper()
        p["email"] = (p["email"] or "").lower()
        p["employee_code"] = p["id"]
        p["full_name"] = " ".join(x for x in (p["first_name"], p["last_name"]) if x)
        p["preferred_name"] = p["first_name"]
        p["cost_centre"] = cost_centre(p["country_code"], p["department"] or "")
        p["company"] = p["company"] or "TimeZ"
        p["source"] = "hcm_sample"
        people.append(p)

    errors: list[str] = []
    ids = [p["id"] for p in people]
    dupes = [i for i, n in Counter(ids).items() if n > 1]
    if None in ids or dupes:
        errors.append(f"blank or duplicate Employee IDs: {dupes[:10]}")
    emails = Counter(p["email"] for p in people)
    if any(n > 1 for n in emails.values()):
        errors.append("duplicate work emails")
    for p in people:
        if not p["joining_date"]:
            errors.append(f"{p['id']}: no joining date")
    id_set = set(ids)
    unknown = sorted({p["manager_id"] for p in people if p["manager_id"] and p["manager_id"] not in id_set})
    if unknown:
        errors.append(f"manager IDs not in the sheet: {unknown[:10]}")
    self_managed = [p["id"] for p in people if p["manager_id"] == p["id"]]
    if self_managed:
        errors.append(f"people managing themselves: {self_managed}")

    manager_of = {p["id"]: p["manager_id"] for p in people}
    in_cycle: set[str] = set()
    for start in manager_of:
        seen: list[str] = []
        node: str | None = start
        while node and node not in seen and node not in in_cycle:
            seen.append(node)
            node = manager_of.get(node)
        if node in seen:
            in_cycle.update(seen[seen.index(node):])
    if in_cycle:
        errors.append(f"reporting-line cycles involving: {sorted(in_cycle)[:10]}")

    by_country = Counter(p["country_code"] for p in people)
    if dict(by_country) != country_summary:
        errors.append(f"country counts {dict(by_country)} don't match the Country Summary sheet {country_summary}")
    by_dept = Counter(p["department"] for p in people)
    if dict(by_dept) != dept_summary:
        errors.append(f"department counts don't match the Department Summary sheet")

    if errors:
        raise SeedDataError("HCM workbook failed validation: " + "; ".join(errors))

    counts = {
        "employees": len(people),
        "top_level": sum(1 for p in people if not p["manager_id"]),
        **{f"country_{k}": v for k, v in sorted(by_country.items())},
    }
    return people, counts


if __name__ == "__main__":
    _, counts = load_employees()
    print("HCM workbook OK: " + ", ".join(f"{k}={v}" for k, v in counts.items()))
